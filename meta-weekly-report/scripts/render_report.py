"""Resolve evidence tokens in report.json; produce matching Markdown and DOCX.

python render_report.py analysis.json report.json --out output_dir
The agent authors the paragraphs according to references/writing-guide.md.
All changing numbers should be {{alias.current.roas|num}}-style tokens.
"""
import argparse, hashlib, json, re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TOKEN=re.compile(r'\{\{([^{}]+)\}\}')

def render(analysis,report,out):
    index={dim:{x['key']:x for x in xs} for dim,xs in analysis['comparisons'].items()}
    refs={}
    for alias,locator in report['refs'].items():
        dim,key=locator.split('::',1);refs[alias]=index[dim][key]
    evidence=[]
    def substitute(match):
        spec=match.group(1);parts=spec.split('|');lookup=parts[0];fmt=parts[1] if len(parts)>1 else 'text'
        alias,*keys=lookup.split('.');v=refs[alias]
        for k in keys:v=v[k]
        evidence.append(dict(token=spec,source=report['refs'][alias],value=v))
        if v is None: return '新增' if fmt=='change' else '-'
        if fmt=='money':return '$'+format(v,',.2f')
        if fmt=='num':return format(v,'.2f')
        if fmt=='int':return format(v,',.0f')
        if fmt=='pct':return format(v*100,'.2f')+'%'
        if fmt=='abs_pct':return format(abs(v)*100,'.2f')+'%'
        if fmt=='change':return ('增长' if v>0 else '下降' if v<0 else '持平')+(format(abs(v)*100,'.2f')+'%' if v else '')
        if fmt=='share':
            period=keys[0];return format(v/refs['total'][period]['spend']*100,'.2f')+'%'
        if fmt=='gmv_share':
            period=keys[0];return format(v/refs['total'][period]['gmv']*100,'.2f')+'%'
        return str(v)
    blocks=[dict(b,text=TOKEN.sub(substitute,b.get('text',''))) for b in report['blocks']]
    if any('{{' in b['text'] for b in blocks):raise ValueError('Unresolved token')
    out.mkdir(parents=True,exist_ok=True)
    stem='Meta周报_'+analysis['period']['current_start'].replace('-','')+'-'+analysis['period']['current_end'].replace('-','')
    md=[];doc=Document()
    # Reference uses Arial 11 body, 18 weekly title, 16 section, 15 subsection.
    section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8);section.left_margin=Cm(1.9);section.right_margin=Cm(1.9)
    for name,size in [('Normal',11),('Title',18),('Heading 1',16),('Heading 2',15),('Heading 3',11)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=1.18
        if name!='Normal':style.font.bold=True;style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(10)
    for style in doc.styles:
        for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for b in blocks:
        role=b.get('role','p');text=b['text']
        style={'title':'Title','h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}.get(role,'Normal')
        p=doc.add_paragraph(style=style)
        if b.get('page_break'):p.paragraph_format.page_break_before=True
        # **bold** mirrors the reference's bold topic opening in prose.
        for i,part in enumerate(re.split(r'\*\*',text)):
            run=p.add_run(part)
            if i%2:run.bold=True
        if role=='blank':p.paragraph_format.space_after=Pt(7)
        if role=='note':
            for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('595959')
        p.paragraph_format.widow_control=True
        md.append({'title':'# ','h1':'## ','h2':'### ','h3':'#### '}.get(role,'')+text)
    # No historical media, commands, actions, or prose are copied into the new report.
    doc.core_properties.title=stem;doc.core_properties.subject='Meta 周度数据总结';doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(out/(stem+'.docx'))
    (out/(stem+'.md')).write_text('\n\n'.join(md)+'\n',encoding='utf-8')
    return dict(stem=stem,blocks=len(blocks),numeric_evidence=evidence)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');p=argparse.ArgumentParser();p.add_argument('analysis',type=Path);p.add_argument('report',type=Path);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    result=render(json.loads(a.analysis.read_text(encoding='utf-8')),json.loads(a.report.read_text(encoding='utf-8')),a.out)
    (a.report.parent/'report-evidence.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(stem=result['stem'],blocks=result['blocks'],evidence_count=len(result['numeric_evidence'])),ensure_ascii=False))
