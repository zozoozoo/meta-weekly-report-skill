"""Deterministic Meta aggregation. Reads XLSX; writes JSON (no Excel authoring).

Usage: python clean_meta.py input.xlsx --out run_dir [--week-end YYYY-MM-DD]
"""
from __future__ import annotations
import argparse, collections, datetime as dt, hashlib, json, math, re, sys
from pathlib import Path
import openpyxl
from openpyxl.utils.datetime import from_excel

TYPE_SET = {'img','vid','video','pic','slide','acimg','acvid'}
BASE = ['spend','gmv','purchases','reach','impressions','clicks','lpv','atc','checkout']
DEFAULTS = {'S225BS737':'KHAKI','S225BS718':'WHITE','S225TS698':'WHITE','S325JS786':'CREME','S325JS804':'BLACK','S124JS443':'STONE','S124TY863':'BLACK'}
ALIASES = {'ITALIAN STRAW':'ITLN STRAW','PERIWINKLE BLUE':'PERIWNKLBL','HALOGEN BLUE':'HAL BLUE','LIGHT GRAY':'LT GRAY'}
SPU = re.compile(r'^([A-Za-z]{1,2}\d{2,3}[A-Za-z]{1,2}\d{3}[A-Za-z]?)(?=$|[_-])')
STRUCT = ['video-视频','合创','img-图片','ugc-红人素材','ecom-电商图','camp-广告图','catalog-目录','slide-幻灯片','SUM-全店汇总','RMT-再营销','拉新受众','ASC-无受众','非ASC','lal-类似受众','int-兴趣词','ADP','PDP']
TABLES = ['说明_口径','表2_花费结构','表3a_受众大类','表3b_受众明细','表4_国家','表5_落地页','表6_素材ID','表7_图片模板','表8_红人','表9_产品x素材']

def norm(x): return re.sub(r'\s+', ' ', str(x or '').strip()).lower()
def div(a,b): return a/b if b else 0.0
def change(c,p): return (c-p)/p if p else None
def numeric(x):
    try:
        n=float(x); return n if math.isfinite(n) else 0.0
    except (ValueError,TypeError): return 0.0

def split_id(value):
    m=re.search(r'\d{4,}',value)
    if not m: return {'creator':'','type':'','number':'','suffix':'','review':'无数字编号'}
    before,after=value[:m.start()],value[m.end():]
    separators=list(re.finditer(r'[_\- ]',before))
    k=separators[-1].start() if separators else -1
    creator=before[:k].strip(' _-') if k>=0 else ''
    typ=before[k+1:] if k>=0 else before
    suffix=re.sub(r'^[_\- ]','',after).strip()
    review=[]
    if re.search(r'(?:^|[_-])v\d+',suffix,re.I): review.append('版本号需复核')
    if re.search(r'[A-Za-z]{1,2}\d{2,3}[A-Za-z]{1,2}\d{3}',suffix) or len(suffix)>50: review.append('后缀含产品或过长')
    if len(typ)>2: review.append('类型需复核')
    return dict(creator=creator,type=typ,number=m.group(),suffix=suffix,review='；'.join(review))

def id_source(value):
    s=norm(value)
    if 'catalog' in s or s=='np': return '其他'
    m=re.search(r'(?:^|_|-)(ip|h|p|v|u)(?=\d)',s)
    if m: return {'ip':'IG热帖','h':'合创','p':'图片','v':'视频','u':'ugc'}[m.group(1)]
    if re.fullmatch(r'[^\d]+?[_\- ]?\d+',s): return '合创'
    return '其他'

def products(value):
    if norm(value) in ('','np'): return [],[]
    keys=[]; issues=[]
    for token in str(value).split('+'):
        token=token.strip(); m=SPU.match(token)
        if not m:
            # Preserve unrecognized tokens for review; never truncate e.g. S326TS265X.
            keys.append('Unknown'); issues.append(token); continue
        code=m.group(1).upper(); tail=token[m.end():]
        color=tail.lstrip('_-').strip().upper() if tail else DEFAULTS.get(code,'')
        color=ALIASES.get(color,color)
        keys.append(code+('_'+color if color else ''))
    return keys,issues

def template_id(value):
    m=re.search(r'inspo[_\-]?([A-Za-z0-9]+)',value,re.I)
    if m:return 'inspo_'+m.group(1).lower()
    m=re.search(r'(?<![A-Za-z])M(\d+)',value,re.I)
    if m:return 'm'+m.group(1)
    return 'mix' if 'mix' in value.lower() else '(无模板号)'

def parse_name(campaign,adset,ad):
    cp=str(campaign or '').split('^'); ap=str(adset or '').split('^'); parts=str(ad or '').split('^')
    get=lambda a,i: a[i].strip() if len(a)>i else ''
    cid=get(parts,1) or 'np'; typ=''; source=''; landing=''
    ti=next((i for i,p in enumerate(parts) if norm(p) in TYPE_SET),None)
    if ti is not None: typ=norm(parts[ti]); source=norm(get(parts,ti+1)); landing=norm(get(parts,ti+2))
    elif 'catalog' in norm(cid): landing=norm(cid)
    landing=re.sub(r'(?:\s*-\s*广告副本)+\s*$','',landing)
    landing=re.sub(r'^adp\+','adp-',landing)
    aud=norm(get(ap,4)) or 'np'
    cat=next((b for a,b in [('int','int'),('lal','lal'),('rmt','RMT'),('daba','daba')] if aud.startswith(a)), 'np' if aud=='np' else '其他')
    ps,issues=products(get(parts,2))
    return dict(cid=cid,idkey=norm(cid),**split_id(cid),id_source=id_source(cid),ad_type=typ,ad_source=source,landing=landing,
                form='ADP' if landing.startswith('adp') else 'PDP' if landing.startswith('pdp') else 'Catalog' if 'catalog' in landing else '其他',
                audience=aud,audience_category=cat,asc=any(norm(get(cp,i))=='asc' for i in (1,2)),budget=norm(get(cp,2)),
                rmt='dpa' in norm(ad)+' '+norm(campaign),products=ps,product_issues=issues,template=template_id(cid))

def structure(r):
    out=['SUM-全店汇总','RMT-再营销' if r['rmt'] else '拉新受众','ASC-无受众' if r['asc'] else '非ASC']
    if r['ad_type'] in ('vid','video','acvid'):out.append('video-视频')
    if r['ad_type'] in ('img','pic','acimg'):out.append('img-图片')
    if r['ad_type']=='slide':out.append('slide-幻灯片')
    for source,label in [('合创','合创'),('ugc','ugc-红人素材'),('ecom','ecom-电商图'),('camp','camp-广告图')]:
        if r['ad_source']==source:out.append(label)
    if 'catalog' in r['landing']:out.append('catalog-目录')
    for prefix,label in [('lal','lal-类似受众'),('int','int-兴趣词')]:
        if r['audience'].startswith(prefix):out.append(label)
    for prefix in ['adp','pdp']:
        if r['landing'].startswith(prefix):out.append(prefix.upper())
    return out

def metrics(rows):
    sums={k:math.fsum(r[k] for r in rows) for k in BASE}
    for name,a,b in [('roas','gmv','spend'),('aov','gmv','purchases'),('ctr','clicks','impressions'),('atc_rate','atc','lpv'),('cvr','purchases','lpv'),('checkout_rate','checkout','atc'),('payment_rate','purchases','checkout')]:sums[name]=div(sums[a],sums[b])
    for name,a,b in [('cpm','spend','impressions'),('cpc','spend','clicks'),('cpa','spend','purchases'),('lpv_cost','spend','lpv')]:
        sums[name]=div(sums[a],sums[b])*(1000 if name=='cpm' else 1)
    return sums

def full_period(rows):
    """All supplied dates, all IDs. Not lifetime outside this source file."""
    cfg=json.loads((Path(__file__).parents[1]/'references'/'material-thresholds.json').read_text(encoding='utf-8'))
    items=list(aggregate(rows,'idkey').values())
    for m in items:
        reasons=[]
        if m['spend']>=cfg['zero_purchase_spend_usd'] and m['purchases']==0:
            reasons.append(f"累计花费≥${cfg['zero_purchase_spend_usd']:,.0f}且购物次数为0")
        if m['spend']>=cfg['low_roas_spend_usd'] and m['roas']<cfg['minimum_roas']:
            reasons.append(f"累计花费≥${cfg['low_roas_spend_usd']:,.0f}且ROAS<{cfg['minimum_roas']}")
        m['flagged']=bool(reasons);m['status']='不达标' if reasons else '未触发规则'
        m['flag_reason']='；'.join(reasons)
    items.sort(key=lambda m:(not m['flagged'],-m['spend'],m['key']))
    total=metrics(rows)
    for k in BASE:
        assert math.isclose(math.fsum(m[k] for m in items),total[k],abs_tol=1e-7),(k,'full_period')
    return dict(start=min(r['day'] for r in rows).isoformat(),end=max(r['day'] for r in rows).isoformat(),
                row_count=len(rows),rules=cfg,total=total,materials=items,flagged_count=sum(m['flagged'] for m in items))

def aggregate(rows,dim):
    groups=collections.defaultdict(list)
    for r in rows:
        if dim=='structure': keys=structure(r)
        elif dim=='product_material':keys=[p+' | '+r['idkey']+' | '+r['id_source'] for p in r['products']]
        elif dim=='products': keys=r['products']
        elif dim=='creator':keys=[norm(r['creator'])] if r['creator'] else []
        elif dim=='template':keys=[r['template']] if r['id_source']=='图片' else []
        elif dim=='landing':keys=[r['landing']] if r['landing'] else []
        elif dim=='country':keys=[r['country']] if norm(r['country']) not in ('','nan') else []
        elif dim=='creator_material': keys=[norm(r['creator'])+' | '+r['idkey']] if r['creator'] else []
        elif dim=='structure_material':keys=[s+' | '+r['idkey'] for s in structure(r)]
        else:keys=[r[dim]]
        for key in keys:groups[key].append(r)
    if dim=='structure':
        for k in STRUCT:groups.setdefault(k,[])
    result={}
    for key,rs in groups.items():
        d=metrics(rs); d['key']=key; d['label']=key; d['rows']=len(rs)
        if dim in ('idkey','product_material','creator_material','structure_material') and rs:
            names=collections.defaultdict(float)
            for r in rs:names[r['cid']]+=r['spend']
            d['label']=sorted(names,key=lambda n:(-names[n],n))[0]
            d.update(split_id(d['label']));d['id_source']=id_source(d['label']);d['products']=sorted(set(p for r in rs for p in r['products']))
            d['first_day']=min(r['day'] for r in rs).isoformat();d['last_day']=max(r['day'] for r in rs).isoformat()
            if dim=='product_material':d['product']=key.split(' | ')[0]
        if dim=='creator' and rs:
            names=collections.defaultdict(float)
            for r in rs:names[r['creator']]+=r['spend']
            d['label']=sorted(names,key=lambda n:(-names[n],n))[0]
        if dim=='landing' and rs:d['form']=rs[0]['form']
        result[key]=d
    return result

def load(path):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    sheets=[s for s in wb if s.max_row>1]
    if len(sheets)!=1:raise ValueError('源文件必须只有一张非空数据表；请明确数据页后再处理。')
    s=sheets[0]
    if s.max_column<18:raise ValueError('列数少于18，不能按旧格式猜测。')
    raw=iter(s.values);headers=next(raw)
    expected=['账户名称','广告系列名称','广告组名称','广告名称','单日','国家/地区','货币']
    if list(headers[:7])!=expected:raise ValueError('前7列字段或列序变化，请核对映射：'+str(headers[:7]))
    if not str(headers[7]).startswith('已花费金额') or list(headers[8:18])!=['购物转化价值','购物次数','覆盖人数','展示次数','链接点击量','落地页浏览量','加入购物车次数','结账发起次数','报告开始日期','报告结束日期']:
        raise ValueError('指标字段或列序变化，请核对第8至18列映射。')
    rows=[]; missing=collections.Counter(); invalid=collections.Counter(); cache={}; duplicates=collections.Counter()
    for line,v in enumerate(raw,2):
        if not any(x not in (None,'') for x in v):continue
        date=v[4]
        try:
            if isinstance(date,(int,float)):date=from_excel(date,wb.epoch)
            if isinstance(date,dt.datetime):date=date.date()
            if not isinstance(date,dt.date):date=dt.date.fromisoformat(str(date)[:10])
        except (ValueError,TypeError):raise ValueError(f'第{line}行日期无效：{v[4]}')
        nk=tuple(v[1:4]); cache.setdefault(nk,parse_name(*nk)); r=dict(cache[nk])
        r.update(day=date,country=str(v[5] or '').strip(),currency=str(v[6] or '').strip(),account=v[0],line=line)
        for key,value in zip(BASE,v[7:16]):
            r[key]=numeric(value)
            if value in (None,''):missing[key]+=1
            else:
                try:
                    if not math.isfinite(float(value)):invalid[key]+=1
                except (ValueError,TypeError):invalid[key]+=1
        rows.append(r);duplicates[tuple(str(x) for x in v)]+=1
    currencies=sorted(set(r['currency'] for r in rows))
    if len(currencies)!=1 or not currencies[0]:raise ValueError('不同货币或空货币不能直接加总：'+str(currencies))
    if currencies[0]!='USD':raise ValueError('当前报告金额格式与绝对变化筛选线按USD配置；请先显式配置其他币种，不进行隐式换汇。')
    return rows,dict(sheet=s.title,headers=list(headers),missing_as_zero=dict(missing),non_numeric_as_zero=dict(invalid),exact_duplicate_extra_rows=sum(n-1 for n in duplicates.values() if n>1),currency=currencies[0])

def build(path,end=None):
    rows,qa=load(path);dates=sorted(set(r['day'] for r in rows))
    end=dt.date.fromisoformat(end) if end else max(dates)
    start=end-dt.timedelta(days=end.weekday());prev_start=start-dt.timedelta(days=7);prev_end=end-dt.timedelta(days=7)
    cr=[r for r in rows if start<=r['day']<=end];pr=[r for r in rows if prev_start<=r['day']<=prev_end]
    if not cr or not pr:raise ValueError('缺少本期或上期数据，不能生成两周对比周报。')
    missing_dates={name:[(a+dt.timedelta(days=i)).isoformat() for i in range((b-a).days+1) if a+dt.timedelta(days=i) not in dates] for name,a,b in [('current',start,end),('previous',prev_start,prev_end)]}
    if any(missing_dates.values()):raise ValueError('比较区间缺少日期，不能默认为零投放：'+str(missing_dates))
    dims=['structure','audience_category','audience','country','landing','idkey','template','creator','product_material','id_source','products','creator_material','structure_material']
    out=dict(version='1.0',source=dict(filename=path.name,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),row_count=len(rows),first_day=min(dates).isoformat(),last_day=max(dates).isoformat()),
             period=dict(current_start=start.isoformat(),current_end=end.isoformat(),previous_start=prev_start.isoformat(),previous_end=prev_end.isoformat(),complete_week=end.weekday()==6,days=end.weekday()+1),
             currency=qa['currency'],tables=TABLES,comparisons={})
    for dim in dims:
        c,p=aggregate(cr,dim),aggregate(pr,dim);result=[]
        for key in sorted(set(c)|set(p)):
            label=(c.get(key) or p[key])['label']
            cm=c.get(key,dict(metrics([]),key=key,label=label));pm=p.get(key,dict(metrics([]),key=key,label=label))
            delta={k:cm[k]-pm[k] for k in metrics([])};wow={k:change(cm[k],pm[k]) for k in metrics([])}
            result.append(dict(key=key,label=cm['label'],current=cm,previous=pm,delta=delta,wow=wow))
        result.sort(key=lambda x:(-x['current']['spend'],-x['previous']['spend'],x['key']))
        if dim=='structure':result.sort(key=lambda x:STRUCT.index(x['key']))
        if dim=='product_material':
            prod_tot=collections.defaultdict(float)
            for x in result:prod_tot[x['key'].split(' | ')[0]]+=x['current']['spend']
            result.sort(key=lambda x:(-prod_tot[x['key'].split(' | ')[0]],x['key'].split(' | ')[0],-x['current']['purchases'],-x['current']['spend']))
        out['comparisons'][dim]=result
    qa['period_row_counts']={'current':len(cr),'previous':len(pr)}
    qa['unknown_products']=dict(collections.Counter(t for r in cr for t in r['product_issues']))
    qa['unknown_product_spend']=math.fsum(r['spend'] for r in cr if r['product_issues'])
    qa['unparsed_ad_type_spend']=math.fsum(r['spend'] for r in cr if not r['ad_type'])
    qa['nonasc_budget_spend']={k:math.fsum(r['spend'] for r in cr if not r['asc'] and r['budget']==k) for k in sorted(set(r['budget'] for r in cr if not r['asc']))}
    qa['excluded_country']=metrics([r for r in cr if norm(r['country']) in ('','nan')])
    qa['excluded_landing']=metrics([r for r in cr if not r['landing']])
    qa['no_creator']=metrics([r for r in cr if not r['creator']])
    qa['no_spu']=metrics([r for r in cr if not r['products']])
    qa['funnel_ratio_over_100']=[{'key':x['key'],'period':p,'checkout_rate':x[p]['checkout_rate'],'payment_rate':x[p]['payment_rate']} for x in out['comparisons']['idkey'] for p in ('current','previous') if x[p]['checkout_rate']>1 or x[p]['payment_rate']>1]
    qa['checks']=[]
    for name,rr in [('current',cr),('previous',pr)]:
        direct=metrics(rr)
        for dim in ('audience_category','audience','idkey','id_source'):
            for k in BASE:
                actual=math.fsum(x[name][k] for x in out['comparisons'][dim])
                assert math.isclose(actual,direct[k],abs_tol=1e-7), (name,dim,k,actual,direct[k])
        for dim,valid in [('country',lambda r:norm(r['country']) not in ('','nan')),('landing',lambda r:bool(r['landing'])),('creator',lambda r:bool(r['creator']))]:
            for k in ('spend','gmv','purchases'):
                actual=math.fsum(x[name][k] for x in out['comparisons'][dim]);expected=math.fsum(r[k] for r in rr if valid(r))
                assert math.isclose(actual,expected,abs_tol=1e-7),(name,dim,k)
        for k in ('spend','gmv','purchases'):
            actual=math.fsum(x[name][k] for x in out['comparisons']['product_material']);expected=math.fsum(r[k]*len(r['products']) for r in rr)
            assert math.isclose(actual,expected,abs_tol=1e-7),(name,'multi_spu',k)
        qa['checks'].append(name+': 总计、互斥维度、剔除项、多SPU乘数核对一致')
    out['qa']=qa
    out['full_period']=full_period(rows)
    cfg=json.loads((Path(__file__).parents[1]/'references'/'analysis-config.json').read_text(encoding='utf-8'))
    total=metrics(cr)['spend'];threshold=max(cfg['minimum_spend_usd'],total*cfg['minimum_spend_share'])
    signals=[]
    for dim in ['country','landing','idkey','template','creator','products']:
        for x in out['comparisons'][dim]:
            c,p=x['current'],x['previous'];w=x['wow'];signs=[]
            material=max(c['spend'],p['spend'])>=threshold
            stable=min(c['purchases'],p['purchases'])>=cfg['minimum_orders_per_period']
            if w['roas'] is not None and abs(w['roas'])>=cfg['roas_change'] and material:signs.append('效率明显改善' if w['roas']>0 else '效率明显下降')
            if abs(x['delta']['gmv'])>=cfg['gmv_change_usd']:signs.append('GMV增量较大' if x['delta']['gmv']>0 else 'GMV减量较大')
            if p['spend']==0 and c['spend']>0:signs.append('上期无花费')
            if c['spend']==0 and p['spend']>=threshold:signs.append('本期无花费')
            if c['purchases']==0 and c['spend']>=threshold:signs.append('有花费未出单')
            if signs:signals.append(dict(dimension=dim,key=x['key'],signals=signs,small_sample=not stable,material=material,delta_gmv=x['delta']['gmv']))
    out['signals']=signals;out['analysis_config']=cfg;out['effective_minimum_spend']=threshold
    # All available weeks provide optional trajectory evidence; use full weeks only.
    out['weekly_history']={}
    for monday in sorted(set(r['day']-dt.timedelta(days=r['day'].weekday()) for r in rows)):
        sunday=monday+dt.timedelta(days=6)
        if all(monday+dt.timedelta(days=i) in dates for i in range(7)):
            wr=[r for r in rows if monday<=r['day']<=sunday]
            out['weekly_history'][monday.isoformat()]={dim:aggregate(wr,dim) for dim in ('structure','idkey')}
    return out

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    ap=argparse.ArgumentParser();ap.add_argument('input',type=Path);ap.add_argument('--out',type=Path,required=True);ap.add_argument('--week-end');args=ap.parse_args()
    result=build(args.input,args.week_end);args.out.mkdir(parents=True,exist_ok=True)
    (args.out/'analysis.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    (args.out/'validation.json').write_text(json.dumps(result['qa'],ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(period=result['period'],source=result['source'],qa=result['qa']),ensure_ascii=False,indent=2))
